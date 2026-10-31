"""The AI tools through the real HTTP routes, with the model replaced by a stand-in."""
import json

import pytest
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.core import llm
from app.tools import autofill, spaces


@pytest.fixture
def run(client):
    def _run(tool_id, file_ids=(), **params):
        response = client.post(f"/api/tools/{tool_id}", json={"file_ids": list(file_ids), "params": params})
        assert response.status_code == 200, response.text
        return response.json()
    return _run


def model_says(monkeypatch, answer):
    seen = []

    def complete(system, user, json_schema=None, model=None):
        seen.append(system + user)
        return answer

    monkeypatch.setattr(llm, "complete", complete)
    return seen


def text_pdf(tmp_path, name, *pages):
    path = tmp_path / name
    c = canvas.Canvas(str(path))
    for text in pages:
        c.drawString(50, 760, text)
        c.showPage()
    c.save()
    return path


def test_all_my_tools_load(client):
    ids = {t["id"] for t in client.get("/api/tools").json()}
    assert {"autofill", "assistant_qa", "assistant_command", "spaces", "generate"} <= ids
    assert not {"autofill", "assistant_qa", "assistant_command", "spaces", "generate"} & set(client.get("/api/tools/_errors").json())


def test_autofill_from_upload_to_filled_file(run, upload, client, monkeypatch, tmp_path):
    monkeypatch.setattr(autofill, "PROFILE_PATH", tmp_path / "p.json")
    monkeypatch.setattr(autofill, "CACHE_PATH", tmp_path / "c.json")
    assert run("autofill", action="save_profile", profile={"surname": "Testperson", "email": "z@example.com"})["ok"]
    file_id = upload("forms/uscis_g1145.pdf")

    seen = model_says(monkeypatch, {"mappings": [{"id": 0, "key": "surname", "confidence": 0.95},
                                                 {"id": 3, "key": "email", "confidence": 0.9}]})
    mapped = run("autofill", [file_id], action="map")
    assert mapped["ok"] and mapped["data"]["source"] == "model"
    assert "Testperson" not in seen[0]

    rows = mapped["data"]["fields"]
    filled = run("autofill", [file_id], action="fill", signature=mapped["data"]["signature"],
                 mapping={r["field"]: r["key"] for r in rows if r["key"]})
    assert filled["ok"] and filled["outputs"][0]["filename"] == "uscis_g1145_filled.pdf"
    pdf = client.get(f"/api/files/{filled['outputs'][0]['file_id']}")
    out = tmp_path / "out.pdf"
    out.write_bytes(pdf.content)
    values = {k.split(".")[-1]: v.get("/V") for k, v in PdfReader(out).get_fields().items()}
    assert values["LastName[0]"] == "Testperson"

    assert run("autofill", [file_id], action="map")["data"]["source"] == "learned"


def test_autofill_asks_for_a_profile_first(run, upload, monkeypatch, tmp_path):
    monkeypatch.setattr(autofill, "PROFILE_PATH", tmp_path / "none.json")
    result = run("autofill", [upload("forms/uscis_g1145.pdf")], action="map")
    assert result["ok"] is False and "profile" in result["message"]


def test_missing_key_is_a_readable_message(run, upload, monkeypatch, tmp_path):
    monkeypatch.setattr(autofill, "PROFILE_PATH", tmp_path / "p.json")
    monkeypatch.setattr(autofill, "CACHE_PATH", tmp_path / "c.json")
    monkeypatch.delenv("SOCLAAS_API_KEY", raising=False)
    monkeypatch.delenv("LLM_BASE_URL", raising=False)
    monkeypatch.setenv("LLM_MODEL", "x")
    run("autofill", action="save_profile", profile={"surname": "T"})
    result = run("autofill", [upload("forms/uscis_g1145.pdf")], action="map")
    assert result["ok"] is False and "SOCLAAS_API_KEY" in result["message"]


def test_qa_route(run, upload, monkeypatch, tmp_path):
    model_says(monkeypatch, {"found": True, "answer": "30 June [p.1]",
                             "citations": [{"page": 1, "quote": "deadline is 30 June"}]})
    file_id = upload(text_pdf(tmp_path, "a.pdf", "The deadline is 30 June."))
    result = run("assistant_qa", [file_id], question="When?")
    assert result["data"]["all_verified"] is True


def test_command_runs_a_plan_across_person_a_and_person_b_tools(run, upload):
    """merge is Person A's, convert is Person B's. Real tools, real files, no model."""
    ids = [upload("forms/irs_w4.pdf"), upload("forms/irs_w9.pdf")]
    steps = [{"tool_id": "merge", "params": {}}, {"tool_id": "convert", "params": {"target": "txt"}}]
    result = run("assistant_command", ids, action="run", steps=steps)
    assert result["ok"], result["message"]
    assert [s["tool_id"] for s in result["data"]["steps"]] == ["merge", "convert"]
    assert result["outputs"][0]["filename"].endswith(".txt")


def test_command_plan_uses_real_schemas_and_catches_bad_ones(run, upload, monkeypatch):
    seen = model_says(monkeypatch, {"explanation": "ok", "steps": [{"tool_id": "convert", "params": {"target": "mp3"}}]})
    result = run("assistant_command", [upload("forms/irs_w4.pdf")], action="plan", sentence="make an mp3")
    assert result["data"]["problems"]
    assert '"esign"' not in seen[0] and '"redact"' in seen[0]


def test_command_refuses_wrong_file_count_before_running(run, upload):
    result = run("assistant_command", [upload("forms/irs_w4.pdf")], action="run", steps=[{"tool_id": "merge", "params": {}}])
    assert result["ok"] is False and "merge" in result["message"]


def test_space_cites_the_right_file_and_page(run, upload, monkeypatch, tmp_path):
    monkeypatch.setattr(spaces, "SPACES_PATH", tmp_path / "s.json")
    a = upload(text_pdf(tmp_path, "lease.pdf", "Rent is 900 a month.", "Deposit is two months."))
    b = upload(text_pdf(tmp_path, "rules.pdf", "No pets are allowed."))
    space = run("spaces", [a, b], action="create", name="Flat")["data"]["space"]
    assert [f["filename"] for f in space["files"]] == ["lease.pdf", "rules.pdf"]
    assert run("spaces", action="save_notes", space_id=space["id"], notes="Ask landlord")["data"]["space"]["notes"] == "Ask landlord"

    model_says(monkeypatch, {"found": True, "answer": "Two months [lease.pdf, p.2]",
                             "citations": [{"file": "lease.pdf", "page": 2, "quote": "Deposit is two months"}]})
    assert run("spaces", action="ask", space_id=space["id"], question="Deposit?")["data"]["all_verified"] is True

    model_says(monkeypatch, {"found": True, "answer": "Two months [rules.pdf, p.1]",
                             "citations": [{"file": "rules.pdf", "page": 1, "quote": "Deposit is two months"}]})
    assert run("spaces", action="ask", space_id=space["id"], question="Deposit?")["data"]["all_verified"] is False

    model_says(monkeypatch, {"found": False, "answer": "x", "citations": []})
    assert run("spaces", action="ask", space_id=space["id"], question="Wifi?")["message"] == "The document doesn't say."


def test_deck_opens_and_has_the_slides(run, upload, client, monkeypatch, tmp_path):
    model_says(monkeypatch, {"title": "Report", "subtitle": "Summary",
                             "slides": [{"title": "One", "bullets": ["a", "b"], "source_pages": [1]}, {"title": "Two", "bullets": ["c"]}]})
    result = run("generate", [upload(text_pdf(tmp_path, "r.pdf", "Body text."))], kind="deck")
    out = tmp_path / "out.pptx"
    out.write_bytes(client.get(f"/api/files/{result['outputs'][0]['file_id']}").content)
    deck = Presentation(str(out))
    assert len(deck.slides) == 3
    assert deck.slides[1].shapes.title.text == "One"


def test_flyer_stays_on_one_page_even_with_too_much_text(run, upload, client, monkeypatch, tmp_path):
    long = "A very long line of flyer copy that goes on and on. " * 12
    model_says(monkeypatch, {"headline": long, "subheading": long, "bullets": [long] * 6, "call_to_action": "Read it"})
    result = run("generate", [upload(text_pdf(tmp_path, "r.pdf", "Body text."))], kind="flyer")
    out = tmp_path / "flyer.pdf"
    out.write_bytes(client.get(f"/api/files/{result['outputs'][0]['file_id']}").content)
    assert len(PdfReader(out).pages) == 1


def test_scan_without_text_needs_ocr_message(run, upload, monkeypatch, tmp_path):
    seen = model_says(monkeypatch, {})
    blank = tmp_path / "blank.pdf"
    c = canvas.Canvas(str(blank))
    c.showPage()
    c.save()
    result = run("assistant_qa", [upload(blank)], question="Anything?")
    assert seen == []
    assert "OCR" in result["message"] or "text" in json.dumps(result["data"])

from b_fixtures import samples  # noqa: F401

from docx import Document
from openpyxl import load_workbook

from app.pro import convert


def test_pdf_to_word(samples):
    res = convert.run([samples["letter"]], {"target": "docx"}, samples["out"])
    text = "\n".join(p.text for p in Document(res["outputs"][0]["path"]).paragraphs)
    assert "Tan Wei Ming" in text and "Supporting details" in text


def test_pdf_table_to_excel(samples):
    res = convert.run([samples["table"]], {"target": "xlsx"}, samples["out"])
    assert res["data"]["tables"] == 1
    ws = load_workbook(res["outputs"][0]["path"]).worksheets[0]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    assert rows[0] == ["Item", "Qty", "Unit price", "Total"]
    assert rows[2] == ["Paper", "5", "6.00", "30.00"]


def test_pdf_to_png_and_back(samples):
    pngs = convert.run([samples["letter"]], {"target": "png", "dpi": 72}, samples["out"])
    assert len(pngs["outputs"]) == 2
    back = convert.run([o["path"] for o in pngs["outputs"]], {"target": "pdf"}, samples["out"])
    assert back["ok"] and back["data"]["pages"] == 2


def test_office_without_libreoffice_says_so(samples, monkeypatch):
    monkeypatch.setattr(convert.shutil, "which", lambda name: None)
    res = convert.run([samples["letter"]], {"target": "office"}, samples["out"])
    assert not res["ok"] and "LibreOffice" in res["message"]

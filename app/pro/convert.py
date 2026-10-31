"""Convert and export (Acrobat Standard/Pro: "convert, export").

params["target"]:
  docx   PDF -> Word: text paragraphs and embedded images, page by page
  xlsx   PDF -> Excel: each detected table on its own sheet (text sheet if none)
  png    PDF -> one PNG per page
  txt    PDF -> plain text with page markers
  pdf    images (PNG/JPG) -> one PDF, one image per page
  office Word/Excel/PowerPoint -> PDF, only if LibreOffice (soffice) is installed
"""
from __future__ import annotations

import shutil
import subprocess
import time
from io import BytesIO
from pathlib import Path

import pdfplumber
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader

from .pdfio import page_count, render_page
from .result import fail, ok


def to_docx(src: Path, out: Path) -> dict:
    doc = Document()
    reader = PdfReader(str(src))
    images = 0
    with pdfplumber.open(str(src)) as pdf:
        for i, page in enumerate(pdf.pages):
            if i:
                doc.add_page_break()
            for line in (page.extract_text() or "").splitlines():
                doc.add_paragraph(line)
            try:
                for im in reader.pages[i].images:
                    try:
                        doc.add_picture(BytesIO(_as_png(im.data)), width=Inches(5.5))
                        images += 1
                    except Exception:
                        pass
            except Exception:
                pass
    doc.save(out)
    return {"pages": page_count(src), "images": images}


def _as_png(data: bytes) -> bytes:
    buf = BytesIO()
    Image.open(BytesIO(data)).convert("RGB").save(buf, format="PNG")
    return buf.getvalue()


def to_xlsx(src: Path, out: Path) -> dict:
    wb = Workbook()
    wb.remove(wb.active)
    tables = 0
    with pdfplumber.open(str(src)) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            for tno, table in enumerate(page.extract_tables(), start=1):
                ws = wb.create_sheet(f"p{pno}_table{tno}"[:31])
                for row in table:
                    ws.append([("" if c is None else c) for c in row])
                tables += 1
        if not tables:
            ws = wb.create_sheet("text")
            for pno, page in enumerate(pdf.pages, start=1):
                for line in (page.extract_text() or "").splitlines():
                    ws.append([pno, line])
    wb.save(out)
    return {"tables": tables}


def run(inputs: list[Path], params: dict, out_dir: Path) -> dict:
    target = params.get("target")
    if not inputs:
        return fail("Choose at least one file to convert.")
    out_dir = Path(out_dir)
    started = time.time()
    src = Path(inputs[0])

    if target == "docx":
        out = out_dir / f"{src.stem}.docx"
        data = to_docx(src, out)
        return ok(f"Converted {data['pages']} page(s) to Word.", [out], {**data, "seconds": round(time.time() - started, 2)})
    if target == "xlsx":
        out = out_dir / f"{src.stem}.xlsx"
        data = to_xlsx(src, out)
        msg = f"Found {data['tables']} table(s)." if data["tables"] else "No tables found; exported the text instead."
        return ok(msg, [out], data)
    if target == "png":
        dpi = int(params.get("dpi", 150))
        outs = []
        for i in range(page_count(src)):
            p = out_dir / f"{src.stem}-page{i + 1}.png"
            render_page(src, i, dpi).save(p)
            outs.append(p)
        return ok(f"Exported {len(outs)} page image(s).", outs, {"pages": len(outs)})
    if target == "txt":
        out = out_dir / f"{src.stem}.txt"
        with pdfplumber.open(str(src)) as pdf:
            parts = [f"--- Page {i} ---\n{(p.extract_text() or '')}" for i, p in enumerate(pdf.pages, start=1)]
        out.write_text("\n\n".join(parts), encoding="utf-8")
        return ok("Exported the text.", [out], {"pages": len(parts)})
    if target == "pdf":
        imgs = [Image.open(p).convert("RGB") for p in inputs]
        out = out_dir / f"{src.stem}.pdf"
        imgs[0].save(out, save_all=True, append_images=imgs[1:], resolution=150)
        return ok(f"Made a {len(imgs)}-page PDF from image(s).", [out], {"pages": len(imgs)})
    if target == "office":
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            return fail("LibreOffice is not installed, so Office files can't be converted on this machine.")
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)],
                       check=True, capture_output=True, timeout=120)
        out = out_dir / f"{src.stem}.pdf"
        return ok("Converted to PDF with LibreOffice.", [out], {})
    return fail("Pick a target: docx, xlsx, png, txt, pdf (from images) or office.")
